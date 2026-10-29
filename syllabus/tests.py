"""
Tests for the syllabus tool.

Run them with:   python manage.py test syllabus

The test files are built in memory by the helper functions below, so
no sample syllabus files need to be saved in the project.
"""

from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from docx import Document

from .extract import ExtractionError, extract_text

SAMPLE_LINES = [
    "CS 101 Introduction to Programming - Fall 2026",
    "Midterm Exam: October 14, in class",
    "Final Project due December 9 at 11:59 PM",
]


def make_docx(lines, table_rows=None):
    """Build a small .docx file in memory and return it as an upload."""
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    if table_rows:
        table = document.add_table(rows=0, cols=len(table_rows[0]))
        for row_values in table_rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row_values):
                cell.text = value
    buffer = BytesIO()
    document.save(buffer)
    return SimpleUploadedFile("syllabus.docx", buffer.getvalue())


def make_pdf(lines):
    """Build a tiny one-page PDF in memory and return it as an upload.

    A PDF is a set of numbered "objects". This writes the smallest set
    that a PDF reader accepts: a catalog, a page list, one page, its
    text, and a font. Real syllabi are made by Word or Google Docs.
    """
    text_commands = "BT /F1 12 Tf 72 720 Td 16 TL "
    for line in lines:
        safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        text_commands += f"({safe}) Tj T* "
    text_commands += "ET"

    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(text_commands)} >>\nstream\n{text_commands}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    pdf = "%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += f"{number} 0 obj\n{body}\nendobj\n"
    xref_start = len(pdf)
    pdf += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    for offset in offsets:
        pdf += f"{offset:010d} 00000 n \n"
    pdf += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_start}\n%%EOF\n"

    return SimpleUploadedFile("syllabus.pdf", pdf.encode("latin-1"))


class ExtractTextTests(SimpleTestCase):
    def test_reads_pdf(self):
        text = extract_text(make_pdf(SAMPLE_LINES))
        self.assertIn("Midterm Exam: October 14", text)
        self.assertIn("--- Page 1 ---", text)

    def test_reads_docx_paragraphs(self):
        text = extract_text(make_docx(SAMPLE_LINES))
        self.assertIn("Final Project due December 9", text)

    def test_reads_docx_tables(self):
        upload = make_docx(
            SAMPLE_LINES,
            table_rows=[["Week", "Date", "Topic"], ["3", "Sep 15", "Loops"]],
        )
        text = extract_text(upload)
        self.assertIn("3 | Sep 15 | Loops", text)

    def test_file_extension_is_not_case_sensitive(self):
        upload = make_docx(SAMPLE_LINES)
        upload.name = "SYLLABUS.DOCX"
        self.assertIn("Midterm", extract_text(upload))

    def test_rejects_other_file_types(self):
        upload = SimpleUploadedFile("syllabus.txt", b"Midterm on Oct 14")
        with self.assertRaisesMessage(ExtractionError, ".pdf or .docx"):
            extract_text(upload)

    def test_rejects_file_with_almost_no_text(self):
        with self.assertRaisesMessage(ExtractionError, "couldn't find any text"):
            extract_text(make_pdf(["Hi"]))

    def test_rejects_damaged_pdf(self):
        upload = SimpleUploadedFile("syllabus.pdf", b"this is not really a pdf")
        with self.assertRaises(ExtractionError):
            extract_text(upload)

    def test_rejects_damaged_docx(self):
        upload = SimpleUploadedFile("syllabus.docx", b"this is not really a docx")
        with self.assertRaisesMessage(ExtractionError, "Word file couldn't be read"):
            extract_text(upload)


# --- Tests for the AI step (ai.py) -----------------------------------------
#
# These never contact Gemini. mock.patch swaps the real Gemini client for
# a fake one that returns whatever answer the test chooses, so the tests
# are free, fast, and work without an API key or internet connection.

import json
from datetime import date
from unittest import mock

from google.genai import errors

from .ai import AIError, build_prompt, find_events

SEMESTER_START = date(2026, 8, 17)


def fake_gemini_reply(text):
    """Make a fake Gemini client whose answer is the given text."""
    fake_client = mock.MagicMock()
    fake_client.models.generate_content.return_value.text = text
    return fake_client


@mock.patch.dict("os.environ", {"GEMINI_API_KEY": "test-key"})
class FindEventsTests(SimpleTestCase):
    def run_with_reply(self, events):
        reply = json.dumps({"events": events})
        with mock.patch("syllabus.ai.genai.Client", return_value=fake_gemini_reply(reply)):
            return find_events("Midterm on October 14 at 1pm.", SEMESTER_START)

    def test_returns_cleaned_events_sorted_by_date(self):
        events = self.run_with_reply([
            {"title": "Final Project due", "date": "2026-12-09", "start_time": "23:59",
             "end_time": None, "category": "assignment", "details": ""},
            {"title": "  Midterm Exam  ", "date": "2026-10-14", "start_time": "13:00",
             "end_time": "14:15", "category": "exam", "details": "Chapters 1-6"},
        ])
        self.assertEqual([e["title"] for e in events], ["Midterm Exam", "Final Project due"])
        self.assertEqual(events[0]["start_time"], "13:00")
        self.assertEqual(events[0]["end_time"], "14:15")

    def test_drops_impossible_dates_and_fixes_bad_times(self):
        events = self.run_with_reply([
            {"title": "Bad date", "date": "2026-02-30", "start_time": None,
             "end_time": None, "category": "other", "details": ""},
            {"title": "Bad time", "date": "2026-09-01", "start_time": "noon",
             "end_time": None, "category": "reading", "details": ""},
        ])
        self.assertEqual(len(events), 1)
        self.assertEqual(events[0]["title"], "Bad time")
        self.assertIsNone(events[0]["start_time"])

    def test_unexpected_reply_format_gives_friendly_error(self):
        with mock.patch("syllabus.ai.genai.Client", return_value=fake_gemini_reply("not json")):
            with self.assertRaisesMessage(AIError, "expected format"):
                find_events("text", SEMESTER_START)

    def test_rate_limit_gives_friendly_error(self):
        fake_client = mock.MagicMock()
        fake_client.models.generate_content.side_effect = errors.APIError(429, {})
        with mock.patch("syllabus.ai.genai.Client", return_value=fake_client):
            with self.assertRaisesMessage(AIError, "free Gemini limit"):
                find_events("text", SEMESTER_START)

    def test_network_problem_gives_friendly_error(self):
        fake_client = mock.MagicMock()
        fake_client.models.generate_content.side_effect = ConnectionError()
        with mock.patch("syllabus.ai.genai.Client", return_value=fake_client):
            with self.assertRaisesMessage(AIError, "internet connection"):
                find_events("text", SEMESTER_START)


class AISetupTests(SimpleTestCase):
    @mock.patch.dict("os.environ", {"GEMINI_API_KEY": ""})
    def test_missing_api_key_gives_friendly_error(self):
        with self.assertRaisesMessage(AIError, "No Gemini API key"):
            find_events("text", SEMESTER_START)

    def test_prompt_includes_semester_start(self):
        prompt = build_prompt("Midterm Oct 14", SEMESTER_START)
        self.assertIn("Monday, August 17, 2026", prompt)
        self.assertIn("Midterm Oct 14", prompt)


# --- Tests for the upload and review pages (views.py) ----------------------
#
# These pretend to be a logged-in user clicking through the pages.
# find_events is replaced with a fake, so the AI is never contacted.

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse

from events.models import Event

FAKE_EVENTS = [
    {"title": "Midterm Exam", "date": "2026-10-14", "start_time": "13:00",
     "end_time": None, "category": "exam", "details": "Chapters 1-6"},
]


class UploadPageTests(TestCase):
    def setUp(self):
        User.objects.create_user("alice", password="test-pass-123")
        self.client.login(username="alice", password="test-pass-123")

    def upload(self, upload_file, semester_start="2026-08-17"):
        return self.client.post(reverse("syllabus_upload"), {
            "syllabus_file": upload_file,
            "semester_start": semester_start,
        })

    def test_requires_login(self):
        self.client.logout()
        response = self.client.get(reverse("syllabus_upload"))
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("login"), response.url)

    def test_page_shows_form_and_nav_link(self):
        response = self.client.get(reverse("syllabus_upload"))
        self.assertContains(response, 'enctype="multipart/form-data"')
        self.assertContains(response, "First day of classes")
        self.assertContains(response, "Import syllabus")  # nav bar link

    @mock.patch("syllabus.views.find_events", return_value=FAKE_EVENTS)
    def test_successful_upload_saves_results_and_redirects(self, fake_find_events):
        response = self.upload(make_docx(SAMPLE_LINES))

        self.assertRedirects(response, reverse("syllabus_review"))
        saved = self.client.session["syllabus_import"]
        self.assertEqual(saved["filename"], "syllabus.docx")
        self.assertEqual(saved["events"], FAKE_EVENTS)
        # The AI was given the semester start as a real date object.
        self.assertEqual(fake_find_events.call_args.args[1], date(2026, 8, 17))

    def test_wrong_file_type_is_rejected_by_form(self):
        response = self.upload(SimpleUploadedFile("notes.txt", b"Midterm Oct 14"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Please choose a .pdf or .docx file.")

    def test_missing_semester_start_is_rejected(self):
        response = self.upload(make_docx(SAMPLE_LINES), semester_start="")
        self.assertContains(response, "This field is required.")

    def test_unreadable_file_shows_message(self):
        response = self.upload(SimpleUploadedFile("broken.pdf", b"not a pdf"))
        self.assertContains(response, "couldn&#x27;t be read")
        self.assertNotIn("syllabus_import", self.client.session)

    @mock.patch("syllabus.views.find_events", side_effect=AIError("The free Gemini limit has been reached for now."))
    def test_ai_problem_shows_message(self, fake_find_events):
        response = self.upload(make_docx(SAMPLE_LINES))
        self.assertContains(response, "free Gemini limit")
        self.assertNotIn("syllabus_import", self.client.session)


class ReviewPageTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("alice", password="test-pass-123")
        self.client.login(username="alice", password="test-pass-123")

    def save_results(self, events):
        """Put AI results in the session, as the upload page would."""
        session = self.client.session
        session["syllabus_import"] = {
            "filename": "CS101.pdf", "semester_start": "2026-08-17", "events": events,
        }
        session.save()

    def post_rows(self, rows):
        """Submit the review form. Each row is a dict of field values.

        This builds the same field names the browser would send:
        events-0-title, events-1-title, ... plus the management form's
        TOTAL_FORMS and INITIAL_FORMS counts.
        """
        data = {
            "events-TOTAL_FORMS": str(len(rows)),
            "events-INITIAL_FORMS": str(len(rows)),
        }
        for number, row in enumerate(rows):
            for field, value in row.items():
                if field == "include" and not value:
                    continue  # an unticked checkbox sends nothing at all
                data[f"events-{number}-{field}"] = "on" if field == "include" else value
        return self.client.post(reverse("syllabus_review"), data)

    def row(self, **changes):
        """A ticked row for the midterm, with any fields changed."""
        values = {"include": True, "title": "Midterm Exam", "date": "2026-10-14",
                  "start_time": "13:00", "end_time": "", "category": "exam",
                  "details": "Chapters 1-6"}
        values.update(changes)
        return values

    # --- Showing the page ---

    def test_without_upload_redirects_to_upload_page(self):
        response = self.client.get(reverse("syllabus_review"))
        self.assertRedirects(response, reverse("syllabus_upload"))

    def test_shows_editable_rows_filled_in(self):
        self.save_results(FAKE_EVENTS)
        response = self.client.get(reverse("syllabus_review"))
        self.assertContains(response, "Review: CS101.pdf")
        self.assertContains(response, 'value="Midterm Exam"')
        self.assertContains(response, 'value="2026-10-14"')
        self.assertContains(response, 'value="13:00"')
        self.assertContains(response, "Add ticked events to calendar")

    def test_no_events_found_message(self):
        self.save_results([])
        response = self.client.get(reverse("syllabus_review"))
        self.assertContains(response, "didn't find any dated events")

    # --- Importing ---

    def test_imports_only_ticked_rows(self):
        self.save_results(FAKE_EVENTS)
        response = self.post_rows([
            self.row(),
            self.row(include=False, title="Wrong item", date="2026-10-20"),
        ])

        self.assertRedirects(response, reverse("calendar_month", args=[2026, 10]))
        event = Event.objects.get()  # .get() fails unless exactly one exists
        self.assertEqual(event.title, "Midterm Exam")
        self.assertEqual(event.owner, self.user)
        self.assertEqual(event.source, "syllabus")
        self.assertEqual(event.category, "exam")
        self.assertEqual(event.description, "Chapters 1-6")
        self.assertEqual(event.local_start().hour, 13)
        self.assertFalse(event.all_day)
        # The upload is finished, so its results are cleared.
        self.assertNotIn("syllabus_import", self.client.session)

    def test_uses_the_users_edits(self):
        self.save_results(FAKE_EVENTS)
        self.post_rows([self.row(title="Midterm (moved)", date="2026-10-16", end_time="14:15")])
        event = Event.objects.get()
        self.assertEqual(event.title, "Midterm (moved)")
        self.assertEqual(event.local_start().date(), date(2026, 10, 16))
        self.assertEqual(event.local_end().hour, 14)

    def test_blank_time_makes_all_day_event(self):
        self.save_results(FAKE_EVENTS)
        self.post_rows([self.row(start_time="")])
        event = Event.objects.get()
        self.assertTrue(event.all_day)
        self.assertIsNone(event.end)

    def test_success_message_on_calendar(self):
        self.save_results(FAKE_EVENTS)
        response = self.post_rows([self.row()])
        response = self.client.get(response.url)
        self.assertContains(response, "Added 1 event from CS101.pdf.")

    def test_skips_events_already_on_calendar(self):
        for attempt in range(2):
            self.save_results(FAKE_EVENTS)
            response = self.post_rows([self.row()])
        self.assertEqual(Event.objects.count(), 1)
        response = self.client.get(response.url)
        self.assertContains(response, "Skipped 1 already on your calendar.")

    # --- Problems the user must fix ---

    def test_ticked_row_without_title_is_rejected(self):
        self.save_results(FAKE_EVENTS)
        response = self.post_rows([self.row(title="  ")])
        self.assertContains(response, "Give this event a name, or untick it.")
        self.assertEqual(Event.objects.count(), 0)

    def test_end_before_start_is_rejected(self):
        self.save_results(FAKE_EVENTS)
        response = self.post_rows([self.row(start_time="14:00", end_time="13:00")])
        self.assertContains(response, "can&#x27;t be before the start time")
        self.assertEqual(Event.objects.count(), 0)

    def test_unticked_blank_row_does_not_block_import(self):
        self.save_results(FAKE_EVENTS)
        self.post_rows([self.row(), self.row(include=False, title="", date="")])
        self.assertEqual(Event.objects.count(), 1)

    def test_nothing_ticked_shows_message(self):
        self.save_results(FAKE_EVENTS)
        response = self.post_rows([self.row(include=False)])
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Tick at least one event")
        self.assertEqual(Event.objects.count(), 0)
        self.assertIn("syllabus_import", self.client.session)  # still there to try again
